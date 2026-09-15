import csv
import io
import unicodedata
from pathlib import Path
from django.core.validators import validate_email
from django.core.exceptions import ValidationError as DjangoValidationError
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Count, F, Min, Prefetch, Q, Sum
from django.utils import timezone
from django.shortcuts import get_object_or_404
from rest_framework import status
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework.exceptions import Throttled
from rest_framework.throttling import AnonRateThrottle
from rest_framework.views import APIView
from .models import AlumnoHabilitado, Area, Asistente, Carrera, ConfiguracionEvento, RetiroCompleto
from .serializers import AreaCatalogoSerializer, AsistenteSerializer, ConfiguracionEventoSerializer, PaseSerializer, RetiroSerializer
from .services import registrar_retiro
from .validators import normalizar_rut, validar_rut

class ExplainedAnonRateThrottle(AnonRateThrottle):
    message = "Se alcanzó el límite de solicitudes desde esta red. Intenta nuevamente en unos minutos."
    def allow_request(self, request, view):
        if super().allow_request(request, view): return True
        raise Throttled(wait=self.wait(), detail=self.message)
class RegistrationThrottle(ExplainedAnonRateThrottle): scope="registration"
class PassRecoveryThrottle(ExplainedAnonRateThrottle): scope="pass_recovery"
class PadronThrottle(ExplainedAnonRateThrottle): scope="padron"
class CatalogoAreasView(APIView):
    permission_classes=[AllowAny]
    throttle_classes=[]
    def get(self,request):
        qs=Area.objects.filter(activa=True).prefetch_related(
            Prefetch("carreras",queryset=Carrera.objects.filter(activa=True),to_attr="carreras_activas")
        )
        return Response(AreaCatalogoSerializer(qs,many=True).data)
class RegistroView(APIView):
    permission_classes=[AllowAny]; throttle_classes=[RegistrationThrottle]
    def post(self,request):
        try: rut=normalizar_rut(request.data.get("rut",""))
        except DjangoValidationError: rut=request.data.get("rut","")
        existing=Asistente.objects.filter(rut=rut).first()
        if existing:
            supplied_email=str(request.data.get("email","")).strip()
            if not supplied_email or supplied_email.casefold() != existing.email.strip().casefold():
                return Response({"detail":"Ese RUT ya está registrado. Si eres tú, ingresa el mismo correo que usaste."},status=status.HTTP_409_CONFLICT)
            return Response(AsistenteSerializer(existing).data|{"recuperado":True})
        s=AsistenteSerializer(data=request.data); s.is_valid(raise_exception=True)
        with transaction.atomic():
            configuracion=ConfiguracionEvento.objects.select_for_update().get_or_create(pk=1)[0]
            if not configuracion.registro_abierto or Asistente.objects.count() >= configuracion.cupo_asistentes:
                detail=configuracion.mensaje_cupos_agotados or "Los cupos para el evento se agotaron."
                return Response({"detail":detail},status=status.HTTP_409_CONFLICT)
            if configuracion.registro_restringido and s.validated_data.get("tipo") == "estudiante":
                alumno=AlumnoHabilitado.objects.select_related("carrera__area").filter(rut=s.validated_data["rut"]).first()
                areas=set(configuracion.areas_prioritarias.values_list("pk",flat=True))
                if alumno is None or (areas and (alumno.carrera_id is None or alumno.carrera.area_id not in areas)):
                    return Response({"detail":MENSAJE_REGISTRO_RESTRINGIDO},status=status.HTTP_403_FORBIDDEN)
            asistente=s.save(completos_asignados=configuracion.completos_por_asistente)
        return Response(AsistenteSerializer(asistente).data,status=status.HTTP_201_CREATED)

def _datos_configuracion(configuracion, incluir_nombres_areas=False):
    registrados=Asistente.objects.count()
    areas=configuracion.areas_prioritarias.order_by("orden","nombre")
    if incluir_nombres_areas:
        areas_prioritarias=list(areas.values("slug","nombre"))
    else:
        areas_prioritarias=list(areas.values_list("slug",flat=True))
    return {"registro_abierto":configuracion.registro_abierto,"cupo_asistentes":configuracion.cupo_asistentes,
            "registrados":registrados,"cupos_disponibles":max(0,configuracion.cupo_asistentes-registrados),
            "completos_por_asistente":configuracion.completos_por_asistente,
            "mensaje_cupos_agotados":configuracion.mensaje_cupos_agotados,
            "registro_restringido":configuracion.registro_restringido,
            "areas_prioritarias":areas_prioritarias}

MENSAJE_REGISTRO_RESTRINGIDO = "Las inscripciones están abiertas por ahora a un grupo de carreras y se abrirán al resto más adelante."

class PadronPrellenadoView(APIView):
    permission_classes=[AllowAny]; throttle_classes=[PadronThrottle]
    def get(self,request,rut):
        configuracion=ConfiguracionEvento.obtener()
        if not configuracion.registro_restringido: return Response(status=status.HTTP_404_NOT_FOUND)
        try: rut=normalizar_rut(rut)
        except DjangoValidationError: return Response(status=status.HTTP_404_NOT_FOUND)
        alumno=AlumnoHabilitado.objects.select_related("carrera__area").filter(rut=rut).first()
        if alumno is None: return Response(status=status.HTTP_404_NOT_FOUND)
        return Response({"nombre":alumno.nombre,"apellido":alumno.apellido,
            "carrera":alumno.carrera.slug if alumno.carrera else None,
            "area":alumno.carrera.area.slug if alumno.carrera else None})

class ConfiguracionEventoView(APIView):
    permission_classes=[AllowAny]
    def get(self,request): return Response(_datos_configuracion(ConfiguracionEvento.obtener(),incluir_nombres_areas=True))

class AdminConfiguracionEventoView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request): return Response(_datos_configuracion(ConfiguracionEvento.obtener()))
    @transaction.atomic
    def patch(self,request):
        configuracion=ConfiguracionEvento.objects.select_for_update().get_or_create(pk=1)[0]
        serializer=ConfiguracionEventoSerializer(configuracion,data=request.data,partial=True)
        serializer.is_valid(raise_exception=True)
        aplicar=serializer.validated_data.pop("aplicar_a_existentes",False)
        cambia_completos="completos_por_asistente" in serializer.validated_data
        configuracion=serializer.save()
        sin_actualizar=0
        if cambia_completos and aplicar:
            nuevo=configuracion.completos_por_asistente
            sin_actualizar=Asistente.objects.filter(completos_retirados__gt=nuevo).count()
            Asistente.objects.filter(completos_retirados__lte=nuevo).update(completos_asignados=nuevo)
        data=_datos_configuracion(configuracion)
        if cambia_completos and aplicar: data["sin_actualizar"]=sin_actualizar
        return Response(data)
class PaseView(APIView):
    permission_classes=[AllowAny]
    def get(self,request,codigo): return Response(PaseSerializer(get_object_or_404(Asistente,codigo=codigo.upper())).data)
class RecuperarPaseView(APIView):
    permission_classes=[AllowAny]; throttle_classes=[PassRecoveryThrottle]
    def post(self,request):
        detail="No se encontró un registro con esos datos."
        try: rut=normalizar_rut(request.data.get("rut",""))
        except DjangoValidationError: return Response({"detail":detail},status=status.HTTP_404_NOT_FOUND)
        asistente=Asistente.objects.filter(rut=rut).first(); email=str(request.data.get("email","")).strip()
        if not asistente or not email or asistente.email.strip().casefold()!=email.casefold():
            return Response({"detail":detail},status=status.HTTP_404_NOT_FOUND)
        return Response({"codigo":asistente.codigo})
class AdminAsistentesView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        qs=Asistente.objects.all().order_by("-creado_en"); q=request.query_params.get("q")
        if q: qs=qs.filter(Q(nombre__icontains=q)|Q(apellido__icontains=q)|Q(rut__icontains=q)|Q(codigo__icontains=q))
        if request.query_params.get("con_saldo")=="1": qs=qs.filter(completos_retirados__lt=F("completos_asignados"))
        return Response(AsistenteSerializer(qs,many=True).data)
class AdminAsistentesExportView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        response=HttpResponse(content_type="text/csv; charset=utf-8"); response["Content-Disposition"]='attachment; filename="asistentes.csv"'; response.write("\ufeff")
        response.write("# Contiene datos personales. La copia impresa debe destruirse al cierre de la jornada.\r\n")
        writer=csv.writer(response); writer.writerow(["código","nombre","apellido","RUT","tipo","área","carrera","completos asignados","completos retirados"])
        for a in Asistente.objects.select_related("area","carrera").order_by("apellido","nombre"):
            writer.writerow([a.codigo,a.nombre,a.apellido,a.rut,a.tipo,a.area.nombre if a.area else "",a.carrera.nombre if a.carrera else "",a.completos_asignados,a.completos_retirados])
        return response
class AdminAsistenteView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request,codigo): return Response(AsistenteSerializer(get_object_or_404(Asistente,codigo=codigo.upper())).data)
class AdminRetirosView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        qs=RetiroCompleto.objects.select_related("asistente","validado_por").all()[:200]
        return Response(RetiroSerializer(qs,many=True).data)
    def post(self,request):
        attendee=get_object_or_404(Asistente,codigo=str(request.data.get("codigo","")).upper())
        try: cantidad=int(request.data.get("cantidad"))
        except (TypeError,ValueError):
            return Response({"detail":"La cantidad debe ser 1 o 2."},status=status.HTTP_400_BAD_REQUEST)
        try: retiro=registrar_retiro(attendee,cantidad,request.user)
        except DjangoValidationError as exc:
            return Response({"detail":exc.messages[0]},status=status.HTTP_400_BAD_REQUEST)
        return Response(RetiroSerializer(retiro).data,status=201)
class AdminResumenView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        from torneos.models import Torneo
        areas=Area.objects.annotate(total=Count("asistente")).filter(total__gt=0).order_by("-total","orden","nombre")
        carreras=Carrera.objects.annotate(total=Count("asistente")).filter(total__gt=0).select_related("area").order_by("-total","nombre")
        configuracion=ConfiguracionEvento.obtener(); registrados=Asistente.objects.count()
        return Response({"total_registrados":registrados,"cupo_asistentes":configuracion.cupo_asistentes,"cupos_disponibles":max(0,configuracion.cupo_asistentes-registrados),"registro_abierto":configuracion.registro_abierto,"registro_restringido":configuracion.registro_restringido,"total_padron":AlumnoHabilitado.objects.count(),"areas_prioritarias":list(configuracion.areas_prioritarias.values_list("slug",flat=True)),"registrados_del_padron":Asistente.objects.filter(rut__in=AlumnoHabilitado.objects.values("rut")).count(),"completos_comprometidos":registrados*configuracion.completos_por_asistente,"completos_entregados":RetiroCompleto.objects.aggregate(v=Sum("cantidad"))["v"] or 0,"completos_pendientes":Asistente.objects.aggregate(v=Sum(F("completos_asignados")-F("completos_retirados")))["v"] or 0,"inscritos_por_torneo":[{"slug":t.slug,"nombre":t.nombre,"inscritos":t.equipos.filter(estado="confirmado").count()} for t in Torneo.objects.all()],"aportes_comprometidos":{x["aporte"]:x["total"] for x in Asistente.objects.values("aporte").annotate(total=Count("id"))},"asistentes_por_area":[{"slug":x.slug,"nombre":x.nombre,"total":x.total} for x in areas],"asistentes_por_carrera":[{"slug":x.slug,"nombre":x.nombre,"area":x.area.slug,"total":x.total} for x in carreras]})

def _sin_tildes(value):
    return "".join(c for c in unicodedata.normalize("NFKD",str(value or "")) if not unicodedata.combining(c)).strip().casefold()

class AdminPadronPlantillaView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb=Workbook(); ws=wb.active; ws.title="Alumnos"
        headers=["rut","nombre","apellido","correo","carrera"]
        ws.append(headers); ws.append(["12.345.678-5","Ana","Ejemplo","ana.ejemplo@example.test","Nombre exacto de carrera"]); ws.append(["11.111.111-1","Luis","Ficticio","","Otra carrera válida"])
        for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="305496")
        for col,width in zip("ABCDE",(18,22,22,32,45)): ws.column_dimensions[col].width=width
        catalogo=wb.create_sheet("Carreras válidas"); catalogo.append(["área","carrera"])
        for cell in catalogo[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="305496")
        for carrera in Carrera.objects.select_related("area").order_by("area__orden","nombre"): catalogo.append([carrera.area.nombre,carrera.nombre])
        catalogo.column_dimensions["A"].width=48; catalogo.column_dimensions["B"].width=48
        output=io.BytesIO(); wb.save(output)
        response=HttpResponse(output.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"]='attachment; filename="plantilla-padron.xlsx"'
        return response

class AdminPadronCargarView(APIView):
    permission_classes=[IsAuthenticated]
    def post(self,request):
        from openpyxl import load_workbook
        archivo=request.FILES.get("archivo"); modo=request.data.get("modo")
        if not archivo or modo not in ("reemplazar","agregar"):
            return Response({"detail":"Debes enviar un archivo .xlsx y un modo válido."},status=400)
        if Path(archivo.name).suffix.casefold() != ".xlsx": return Response({"detail":"El archivo debe tener extensión .xlsx."},status=400)
        try: wb=load_workbook(archivo,read_only=True,data_only=True)
        except Exception: return Response({"detail":"No se pudo leer el archivo Excel."},status=400)
        ws=wb["Alumnos"] if "Alumnos" in wb.sheetnames else wb.worksheets[0]
        rows=ws.iter_rows(values_only=True)
        try: raw_headers=next(rows)
        except StopIteration: return Response({"detail":"El archivo está vacío."},status=400)
        aliases={"rut":"rut","nombre":"nombre","apellido":"apellido","correo":"email","email":"email","carrera":"carrera"}
        columns={aliases[h]:i for i,value in enumerate(raw_headers) if (h:=_sin_tildes(value)) in aliases}
        missing={"rut","nombre","apellido","carrera"}-set(columns)
        if missing: return Response({"detail":"Faltan encabezados obligatorios: "+", ".join(sorted(missing))},status=400)
        carreras=list(Carrera.objects.select_related("area")); by_name={_sin_tildes(c.nombre):c for c in carreras}; by_slug={c.slug.casefold():c for c in carreras}
        staged={}; errors=[]; total_errors=0; omitted=0
        for row_number,row in enumerate(rows,start=2):
            if not any(value not in (None,"") for value in row): continue
            def value(key):
                index=columns.get(key); return str(row[index] or "").strip() if index is not None and index < len(row) else ""
            raw_rut=value("rut")
            try: rut=validar_rut(raw_rut)
            except DjangoValidationError as exc:
                total_errors+=1; errors.append({"fila":row_number,"rut":raw_rut,"motivo":exc.messages[0]}); continue
            carrera_text=value("carrera"); carrera=by_name.get(_sin_tildes(carrera_text)) or by_slug.get(carrera_text.casefold())
            motivo=None
            if not value("nombre") or not value("apellido"): motivo="Nombre y apellido son obligatorios."
            elif carrera is None: motivo="Carrera no válida."
            email=value("email")
            if not motivo and email:
                try: validate_email(email)
                except DjangoValidationError: motivo="Correo no válido."
            if motivo:
                total_errors+=1; errors.append({"fila":row_number,"rut":raw_rut,"motivo":motivo}); continue
            if rut in staged: omitted+=1
            staged[rut]={"nombre":value("nombre"),"apellido":value("apellido"),"email":email,"carrera":carrera}
        errors=errors[:50]
        if total_errors or not staged:
            data={"importados":0,"actualizados":0,"omitidos":omitted,"total_padron":AlumnoHabilitado.objects.count(),"errores":errors,"total_errores":total_errors}
            if not staged and not total_errors: data["detail"]="El archivo no contiene filas válidas."
            return Response(data,status=400)
        lote=f'{timezone.localtime():%Y-%m-%d %H:%M} {Path(archivo.name).name}'[:60]
        with transaction.atomic():
            previous=set(AlumnoHabilitado.objects.filter(rut__in=staged).values_list("rut",flat=True)) if modo=="agregar" else set()
            if modo=="reemplazar": AlumnoHabilitado.objects.all().delete()
            for rut,data in staged.items(): AlumnoHabilitado.objects.update_or_create(rut=rut,defaults=data|{"lote":lote})
        return Response({"importados":len(staged)-len(previous),"actualizados":len(previous),"omitidos":omitted,"total_padron":AlumnoHabilitado.objects.count(),"errores":[],"total_errores":0})

class AdminPadronView(APIView):
    permission_classes=[IsAuthenticated]
    def get(self,request):
        areas=AlumnoHabilitado.objects.values("carrera__area__slug","carrera__area__nombre").annotate(total=Count("id")).order_by("carrera__area__nombre")
        carreras=AlumnoHabilitado.objects.values("carrera__slug","carrera__nombre","carrera__area__slug").annotate(total=Count("id")).order_by("carrera__nombre")
        lotes=AlumnoHabilitado.objects.values("lote").annotate(fecha=Min("creado_en"),cantidad=Count("id")).order_by("-fecha")
        return Response({"total":AlumnoHabilitado.objects.count(),"por_area":[{"slug":x["carrera__area__slug"],"nombre":x["carrera__area__nombre"],"total":x["total"]} for x in areas],"por_carrera":[{"slug":x["carrera__slug"],"nombre":x["carrera__nombre"],"area":x["carrera__area__slug"],"total":x["total"]} for x in carreras],"lotes":[{"lote":x["lote"],"fecha":x["fecha"],"cantidad":x["cantidad"]} for x in lotes]})
    def delete(self,request):
        if request.data.get("confirmacion") != "BORRAR PADRON": return Response({"detail":"Confirmación inválida."},status=400)
        if ConfiguracionEvento.obtener().registro_restringido: return Response({"detail":"Desactiva el registro restringido antes de borrar el padrón."},status=400)
        deleted=AlumnoHabilitado.objects.count(); AlumnoHabilitado.objects.all().delete()
        return Response({"eliminados":deleted,"total_padron":0})
